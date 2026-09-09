import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Brand Colors
    C_BG = RGBColor(247, 246, 242)        # #F7F6F2
    C_DARK = RGBColor(24, 24, 22)         # #181816
    C_BLUE = RGBColor(26, 86, 219)        # #1A56DB
    C_WHITE = RGBColor(255, 255, 255)
    C_GRAY = RGBColor(110, 110, 105)
    C_LIGHT_GRAY = RGBColor(230, 228, 222)
    C_CARD = RGBColor(255, 255, 255)
    C_BLUE_BG = RGBColor(238, 242, 255)

    def set_slide_background(slide, color=C_BG):
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = color
        bg_shape.line.fill.background() # no line

    def add_card(slide, left, top, width, height, fill_color=C_CARD, line_color=C_LIGHT_GRAY):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        if line_color:
            shape.line.color.rgb = line_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    def add_header(slide, category, title, subtitle=None):
        # Category tag
        tx_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.55), Inches(11.5), Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_BLUE
        p.font.name = "Arial"

        # Title
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = C_DARK
        p2.font.name = "Malgun Gothic"
        p2.space_before = Pt(4)

        if subtitle:
            p3 = tf.add_paragraph()
            p3.text = subtitle
            p3.font.size = Pt(12)
            p3.font.color.rgb = C_GRAY
            p3.font.name = "Malgun Gothic"
            p3.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Cover)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, C_BG)

    # Main Card container
    add_card(s1, Inches(0.9), Inches(0.9), Inches(11.533), Inches(5.7), C_WHITE, None)

    # Top Brand Pill
    pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(2.2), Inches(0.45))
    pill.fill.solid()
    pill.fill.fore_color.rgb = C_BLUE_BG
    pill.line.fill.background()
    p_tf = pill.text_frame
    p_p = p_tf.paragraphs[0]
    p_p.text = "ART CURATION PLATFORM"
    p_p.alignment = PP_ALIGN.CENTER
    p_p.font.size = Pt(9.5)
    p_p.font.bold = True
    p_p.font.color.rgb = C_BLUE

    # Main Title
    tb1 = s1.shapes.add_textbox(Inches(1.5), Inches(2.1), Inches(8.5), Inches(2.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "ARTPICK"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p.font.name = "Georgia"

    p2 = tf1.add_paragraph()
    p2.text = "당신의 취향이 작품을 만나는 곳"
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = C_BLUE
    p2.font.name = "Malgun Gothic"
    p2.space_before = Pt(10)

    p3 = tf1.add_paragraph()
    p3.text = "새로운 작가와 작품을 발견하고 소장하는 차세대 아트 큐레이션 플랫폼"
    p3.font.size = Pt(14)
    p3.font.color.rgb = C_GRAY
    p3.font.name = "Malgun Gothic"
    p3.space_before = Pt(12)

    # Bottom Metadata
    tb_foot = s1.shapes.add_textbox(Inches(1.5), Inches(5.3), Inches(10.0), Inches(0.8))
    tf_foot = tb_foot.text_frame
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = "2026.09  |  ARTPICK Service Launch Presentation  |  ARTISTS · ARTWORKS · COMMUNITY"
    p_foot.font.size = Pt(11)
    p_foot.font.bold = True
    p_foot.font.color.rgb = C_GRAY

    # -------------------------------------------------------------
    # SLIDE 2: Background & Problem (기획 배경 및 문제 정의)
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, C_BG)
    add_header(s2, "Background & Problem", "기존 미술 시장의 한계와 ARTPICK의 출발점", "누구나 미술을 쉽게 발견하고 소장할 수 있도록 문턱을 낮춥니다.")

    # 3 Problem Cards
    cards_data_s2 = [
        ("01. 높은 진입 장벽", "정보 비대칭성과 폐쇄적 시장", "갤러리와 옥션 중심의 유통 구조로 인해 가격이 불투명하고 초보 컬렉터의 진입 장벽이 높았습니다."),
        ("02. 작가의 노출 기회 부족", "신진 아티스트의 고립", "매년 수천 명의 유망 작가가 배출되지만, 오프라인 갤러리 전시 기회는 극소수에게만 주어졌습니다."),
        ("03. 파편화된 취향 탐색", "개인 맞춤형 큐레이션 부재", "방대한 미술 작품 속에서 자신의 취향, 인테리어 분위기에 어울리는 원화를 직관적으로 찾기 어려웠습니다.")
    ]
    for i, (num, title, desc) in enumerate(cards_data_s2):
        x = Inches(0.9 + i * 3.9)
        card = add_card(s2, x, Inches(1.8), Inches(3.733), Inches(4.7), C_WHITE, C_LIGHT_GRAY)
        tb = s2.shapes.add_textbox(x + Inches(0.3), Inches(2.2), Inches(3.133), Inches(3.9))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = num
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_BLUE
        
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = C_DARK
        p2.font.name = "Malgun Gothic"
        p2.space_before = Pt(8)
        
        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(12.5)
        p3.font.color.rgb = C_GRAY
        p3.font.name = "Malgun Gothic"
        p3.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 3: Brand Identity & Philosophy (브랜드 아이덴티티)
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, C_BG)
    add_header(s3, "Brand Identity & Value", "ARTPICK의 철학: 예술이 사람을 더 가깝게 만듭니다", "공간을 바꾸는 작은 선택이 예술 생태계의 내일을 만듭니다.")

    # Left Big Card
    add_card(s3, Inches(0.9), Inches(1.8), Inches(5.6), Inches(4.7), C_DARK, None)
    tb_l = s3.shapes.add_textbox(Inches(1.3), Inches(2.3), Inches(4.8), Inches(3.8))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    
    p = tf_l.paragraphs[0]
    p.text = "CORE PHILOSOPHY"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(96, 165, 250)
    
    p2 = tf_l.add_paragraph()
    p2.text = "“예술을 고르는 당신이,\n새로운 이야기를 만듭니다.”"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = C_WHITE
    p2.font.name = "Malgun Gothic"
    p2.space_before = Pt(12)
    
    p3 = tf_l.add_paragraph()
    p3.text = "Good Art Brings People Together.\n원화를 소장하는 특별한 설렘과 작가의 진정성 있는 스토리를 온전히 연결합니다."
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGBColor(200, 200, 200)
    p3.font.name = "Malgun Gothic"
    p3.space_before = Pt(16)

    # Right: 4 Pillars
    pillars = [
        ("ARTISTS", "창작에만 몰입할 수 있는 투명한 유통 생태계"),
        ("ARTWORKS", "회화·사진·일러스트·조각을 아우르는 엄선된 컬렉션"),
        ("COMMUNITY", "작가와 컬렉터가 소통하고 전시를 함께하는 커뮤니티"),
        ("A BRIGHTER TOMORROW", "작은 소장이 모여 예술의 내일을 밝히는 선순환")
    ]
    for j, (p_title, p_desc) in enumerate(pillars):
        y = Inches(1.8 + j * 1.2)
        add_card(s3, Inches(6.8), y, Inches(5.633), Inches(1.05), C_WHITE, C_LIGHT_GRAY)
        tb_p = s3.shapes.add_textbox(Inches(7.1), y + Inches(0.12), Inches(5.1), Inches(0.8))
        tf_p = tb_p.text_frame
        tf_p.word_wrap = True
        
        p = tf_p.paragraphs[0]
        p.text = p_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_BLUE
        
        p2 = tf_p.add_paragraph()
        p2.text = p_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_DARK
        p2.font.name = "Malgun Gothic"
        p2.space_before = Pt(2)

    # -------------------------------------------------------------
    # SLIDE 4: Desktop Web Experience (핵심 기능: 데스크톱 웹)
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, C_BG)
    add_header(s4, "Web Service Features", "데스크톱 웹: 미술관 도록 수준의 우아한 큐레이션", "고해상도 비주얼 중심 UI와 직관적인 카테고리 인터랙션을 제공합니다.")

    features_s4 = [
        ("01. 히어로 갤러리 & 지표", "1,200+ 작가, 8,500+ 작품", "대형 블루 추상화 중심의 미술관 감성 레이아웃과 실시간 플랫폼 성장 수치를 제공하여 브랜드 신뢰도를 극대화합니다."),
        ("02. 주목할 만한 작품 피드", "회화·사진·일러스트·조각 필터", "원클릭 카테고리 필터링과 비대칭 갤러리 그리드로 작품의 매력을 시각적으로 극대화하며, 즉시 찜(좋아요) 저장이 가능합니다."),
        ("03. 작품 상세 및 소장 문의", "투명한 규격·기법 및 보증서", "호수, 크기, 재료, 제작연도, 친필 서명 진품 보증서를 투명하게 안내하고 원클릭 소장/구매 상담 시스템을 연결합니다."),
        ("04. 실시간 작품 등록 스튜디오", "창작자를 위한 직관적 업로드", "신진 작가가 자신의 작품과 작가 노트를 손쉽게 등록하여 즉시 플랫폼에 노출하고 컬렉터와 매칭될 수 있습니다.")
    ]
    for k, (f_title, f_sub, f_desc) in enumerate(features_s4):
        x = Inches(0.9 + (k % 2) * 5.9)
        y = Inches(1.8 + (k // 2) * 2.5)
        add_card(s4, x, y, Inches(5.633), Inches(2.25), C_WHITE, C_LIGHT_GRAY)
        tb = s4.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), Inches(5.0), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f_title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_BLUE
        p.font.name = "Malgun Gothic"
        
        p2 = tf.add_paragraph()
        p2.text = f_sub
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = C_DARK
        p2.font.name = "Malgun Gothic"
        p2.space_before = Pt(2)
        
        p3 = tf.add_paragraph()
        p3.text = f_desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = C_GRAY
        p3.font.name = "Malgun Gothic"
        p3.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 5: Mobile App Experience (모바일 피드 & 작가 허브)
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, C_BG)
    add_header(s5, "Mobile Service Features", "모바일 앱: 손안에서 만나는 피드 & 작가 프로필", "일상 속에서 가볍게 작품을 탐색하고 선호하는 작가를 팔로우합니다.")

    # Left: Screen 1 Explanation
    add_card(s5, Inches(0.9), Inches(1.8), Inches(5.6), Inches(4.7), C_WHITE, C_LIGHT_GRAY)
    tb_s5_l = s5.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(5.0), Inches(4.0))
    tf_s5_l = tb_s5_l.text_frame
    tf_s5_l.word_wrap = True
    
    p = tf_s5_l.paragraphs[0]
    p.text = "화면 1: 모바일 홈 피드 탐색"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.font.name = "Malgun Gothic"
    
    items_l = [
        "• 통합 퀵 검색창: 작품명, 작가명, 키워드 실시간 서치",
        "• 원터치 카테고리 필터: 전체, 회화, 사진, 일러스트, 조각",
        "• 2열 비주얼 카드 피드: 최적화된 모바일 감상 경험",
        "• 원터치 하트(좋아요): 마이페이지 찜 보관함 실시간 동기화",
        "• 하단 글로벌 탭바: [홈, 탐색, 등록, 좋아요, 마이] 5개 탭"
    ]
    for item in items_l:
        pi = tf_s5_l.add_paragraph()
        pi.text = item
        pi.font.size = Pt(12)
        pi.font.color.rgb = C_DARK
        pi.font.name = "Malgun Gothic"
        pi.space_before = Pt(8)

    # Right: Screen 2 Explanation
    add_card(s5, Inches(6.833), Inches(1.8), Inches(5.6), Inches(4.7), C_WHITE, C_LIGHT_GRAY)
    tb_s5_r = s5.shapes.add_textbox(Inches(7.133), Inches(2.1), Inches(5.0), Inches(4.0))
    tf_s5_r = tb_s5_r.text_frame
    tf_s5_r.word_wrap = True
    
    p = tf_s5_r.paragraphs[0]
    p.text = "화면 2: 작가 프로필 허브 (김서영 작가)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.font.name = "Malgun Gothic"
    
    items_r = [
        "• 작가 아이덴티티: 원형 아바타, 대표 커버 배너, 핸들(@seoyoung_kim)",
        "• 팬덤 팔로우 시스템: 원클릭 팔로우 토글 및 팔로워 수 실시간 카운트",
        "• 핵심 활동 지표: 1.2만 팔로워 · 24 작품 · 3 전시",
        "• 감성 작가 노트: “평범한 것들이 모여 특별한 하루가 됩니다.”",
        "• 4대 탭 네비게이션: [작품(3x3 썸네일 그리드), 소개, 전시 이력, 소식]"
    ]
    for item in items_r:
        pi = tf_s5_r.add_paragraph()
        pi.text = item
        pi.font.size = Pt(12)
        pi.font.color.rgb = C_DARK
        pi.font.name = "Malgun Gothic"
        pi.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 6: Featured Artworks & Creators (대표 작품 및 작가군)
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, C_BG)
    add_header(s6, "Curated Creators", "다양한 장르를 아우르는 ARTPICK 대표 라인업", "회화부터 사진, 조각, 일러스트까지 검증된 아티스트들이 함께합니다.")

    artworks_data = [
        ("김서영", "푸른 하루", "₩1,200,000", "회화 / 수국 유화 30호"),
        ("이준호", "저녁의 바다", "₩800,000", "사진 / 아날로그 필름 프린트"),
        ("정민지", "그리고, 또 하루", "₩950,000", "일러스트 / 비비드 팝 아트"),
        ("박도현", "흐르는 형태", "₩2,500,000", "조각 / 백자토 미니멀 오브제"),
        ("최은비", "골목의 시간", "₩700,000", "사진 / 서촌 흑백 다큐멘터리"),
        ("한지호", "여름의 결", "₩1,600,000", "회화 / 산자락 대형 유화 50호")
    ]
    for idx, (artist, title, price, spec) in enumerate(artworks_data):
        row = idx // 3
        col = idx % 3
        x = Inches(0.9 + col * 3.9)
        y = Inches(1.8 + row * 2.45)
        add_card(s6, x, y, Inches(3.733), Inches(2.2), C_WHITE, C_LIGHT_GRAY)
        
        tb = s6.shapes.add_textbox(x + Inches(0.25), y + Inches(0.2), Inches(3.2), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = artist
        p.font.size = Pt(12)
        p.font.color.rgb = C_GRAY
        p.font.name = "Malgun Gothic"
        
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = C_DARK
        p2.font.name = "Malgun Gothic"
        p2.space_before = Pt(2)
        
        p3 = tf.add_paragraph()
        p3.text = price
        p3.font.size = Pt(15)
        p3.font.bold = True
        p3.font.color.rgb = C_BLUE
        p3.space_before = Pt(4)
        
        p4 = tf.add_paragraph()
        p4.text = spec
        p4.font.size = Pt(10.5)
        p4.font.color.rgb = C_GRAY
        p4.font.name = "Malgun Gothic"
        p4.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 7: Business Model & Scalability (비즈니스 모델)
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, C_BG)
    add_header(s7, "Business Model", "지속 가능한 아트 생태계와 수익 모델", "작가와 플랫폼, 컬렉터가 함께 성장하는 상생 구조를 지향합니다.")

    bm_data = [
        ("01. 원화 중개 수수료", "15~20%의 합리적 수수료", "기존 오프라인 갤러리(40~50%) 대비 획기적으로 낮춘 작가 친화적 중개 수수료로 양질의 작품을 유치합니다."),
        ("02. 한정판 에디션 & 아트 굿즈", "프린트 및 굿즈 커머스", "원화 소장이 부담스러운 MZ 컬렉터를 위한 리미티드 에디션 판화, 포스터, 아트북 콜라보레이션 상품 판매."),
        ("03. B2B 공간 아트 렌탈", "오피스·호텔 맞춤 큐레이션", "기업 사옥, 프리미엄 카페, 호텔 라운지에 정기적으로 작품을 교체 전시하는 구독형 공간 아트 솔루션 제공."),
        ("04. 에디토리얼 브랜디드 콘텐츠", "매거진 & 스폰서십", "작가 심층 인터뷰, 컬렉팅 가이드 콘텐츠를 기반으로 한 아트 브랜드 제휴 및 프리미엄 프로모션.")
    ]
    for b_idx, (b_title, b_sub, b_desc) in enumerate(bm_data):
        x = Inches(0.9 + (b_idx % 2) * 5.9)
        y = Inches(1.8 + (b_idx // 2) * 2.5)
        add_card(s7, x, y, Inches(5.633), Inches(2.25), C_WHITE, C_LIGHT_GRAY)
        tb = s7.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), Inches(5.0), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = b_title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_BLUE
        p.font.name = "Malgun Gothic"
        
        p2 = tf.add_paragraph()
        p2.text = b_sub
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = C_DARK
        p2.font.name = "Malgun Gothic"
        p2.space_before = Pt(2)
        
        p3 = tf.add_paragraph()
        p3.text = b_desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = C_GRAY
        p3.font.name = "Malgun Gothic"
        p3.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 8: Roadmap & Vision (로드맵 및 비전)
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, C_BG)
    add_header(s8, "Roadmap & Vision", "미래를 향한 로드맵: More Artists, A Brighter Tomorrow", "온·오프라인의 경계를 넘어 국내외 미술 시장의 표준 플랫폼으로 도약합니다.")

    # 3 Timeline Cards
    timeline = [
        ("PHASE 1 (현재 완료)", "플랫폼 런칭 & 생태계 구축", "• 데스크톱 웹 & 모바일 피드 런칭\n• 등록 작가 1,200명 유치 완료\n• 투명한 소장 문의 프로세스 정립\n• 작가 프로필 허브 시스템 구축"),
        ("PHASE 2 (2026.4Q)", "스마트 큐레이션 고도화", "• AI 기반 개인 취향 분석 추천 엔진\n• AR 카메라를 통한 벽면 가상 배치\n• 간편 전자 결제 & 안심 배송 연동\n• 신진 작가 오프라인 기획전 개최"),
        ("PHASE 3 (2027.1H)", "글로벌 확장 & 아트페어", "• K-아트 글로벌 해외 배송 파이프라인\n• 자체 ARTPICK 온·오프라인 아트페어\n• 기업 B2B 아트 구독 서비스 본격화\n• 국내 최대 아트 커뮤니티 정착")
    ]
    for t_idx, (phase, subtitle, bullets) in enumerate(timeline):
        x = Inches(0.9 + t_idx * 3.9)
        add_card(s8, x, Inches(1.8), Inches(3.733), Inches(4.0), C_WHITE, C_LIGHT_GRAY)
        tb = s8.shapes.add_textbox(x + Inches(0.25), Inches(2.05), Inches(3.233), Inches(3.5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = phase
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_BLUE
        
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = C_DARK
        p2.font.name = "Malgun Gothic"
        p2.space_before = Pt(6)
        
        p3 = tf.add_paragraph()
        p3.text = bullets
        p3.font.size = Pt(11.5)
        p3.font.color.rgb = C_GRAY
        p3.font.name = "Malgun Gothic"
        p3.space_before = Pt(12)

    # Bottom Banner
    bot_card = add_card(s8, Inches(0.9), Inches(6.0), Inches(11.533), Inches(0.9), C_DARK, None)
    tb_b = s8.shapes.add_textbox(Inches(1.2), Inches(6.15), Inches(11.0), Inches(0.6))
    tf_b = tb_b.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "ARTPICK — 작은 선택이 새로운 예술의 내일을 만듭니다.  |  Thank You."
    p_b.alignment = PP_ALIGN.CENTER
    p_b.font.size = Pt(14)
    p_b.font.bold = True
    p_b.font.color.rgb = C_WHITE

    # Save to file
    out_path = r"C:\Artpick\ARTPICK_발표자료.pptx"
    prs.save(out_path)
    print(f"Presentation saved to: {out_path}")

if __name__ == "__main__":
    create_presentation()
